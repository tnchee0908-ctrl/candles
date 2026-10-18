"""Файлы между пользователем и песочницей Code Interpreter.

Загруженный файл уходит в контейнер OpenAI, поэтому ПДн в нём быть не должно. Принцип — белый список,
а не «найти и заменить»: из файла берётся только то, что нужно для расчёта (значения ячеек, текст),
каждое значение маскируется, и из этого собирается НОВЫЙ файл. Свойства документа (автор, кто изменил),
комментарии к ячейкам и их авторы, колонтитулы, имена листов-шаблонов и прочие служебные части
в него просто не попадают. Имя файла в контейнере тоже нейтральное: upload_<id>.<ext>.
"""

import csv
import io
import json
import logging
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

from app.assistant.pii import Pseudonymizer
from app.errors import DomainError

log = logging.getLogger(__name__)

TEXT_EXT = {".csv", ".txt", ".tsv", ".json", ".md"}
TABLE_EXT = {".csv": ",", ".tsv": "\t"}
# Защита от «zip-бомбы»: XLSX — это zip, маленький файл может распаковаться в гигабайты.
MAX_XLSX_UNPACKED = 50 * 1024 * 1024
MAX_XLSX_PARTS = 2000
MAX_SHEET_CELLS = 1_000_000


def decode_text(data: bytes) -> str:
    """Текст в любой из типичных кодировок: UTF-8 (с BOM и без), UTF-16 (с BOM и без), cp1251."""
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    # UTF-16 без BOM: каждый второй байт латиницы/цифр — ноль.
    sample = data[:4000]
    if sample and sample.count(0) > len(sample) // 4:
        even_zeros = sample[0::2].count(0)
        odd_zeros = sample[1::2].count(0)
        return data.decode("utf-16-be" if even_zeros > odd_zeros else "utf-16-le", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1251", errors="replace")


def check_xlsx(data: bytes) -> None:
    """Отказ до разбора: не zip, слишком много частей или слишком большой распакованный объём."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile as exc:
        raise DomainError("Файл не похож на XLSX: не удалось его открыть.") from exc
    if len(infos) > MAX_XLSX_PARTS or sum(i.file_size for i in infos) > MAX_XLSX_UNPACKED:
        raise DomainError(
            "XLSX слишком большой после распаковки (больше 50 МБ). Сохраните нужный лист в CSV."
        )


def rebuild_xlsx(data: bytes, fn: Callable[[str], str]) -> bytes:
    """Новая книга только из значений ячеек (строки пропущены через fn). Остальное не переносится."""
    from openpyxl import Workbook, load_workbook

    check_xlsx(data)
    src = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    out = Workbook()
    out.remove(out.active)
    used_titles: set[str] = set()
    for n, ws in enumerate(src.worksheets, start=1):
        title = _sheet_title(fn(ws.title), n, used_titles)
        dst = out.create_sheet(title)
        cells = 0
        for row in ws.iter_rows(values_only=True):
            cells += len(row)
            if cells > MAX_SHEET_CELLS:
                raise DomainError("В листе XLSX слишком много ячеек. Сохраните нужный диапазон в CSV.")
            dst.append([fn(v) if isinstance(v, str) else v for v in row])
    src.close()
    # Свойства новой книги — пустые: openpyxl иначе подставит своего «автора».
    out.properties.creator = None
    out.properties.lastModifiedBy = None
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def _sheet_title(title: str, n: int, used: set[str]) -> str:
    cleaned = "".join(ch for ch in title if ch not in "[]:*?/\\").strip()[:31] or f"Лист{n}"
    while cleaned in used:
        cleaned = f"{cleaned[:28]}_{n}"
    used.add(cleaned)
    return cleaned


def _mask_table(text: str, delimiter: str, fn: Callable[[str], str]) -> str:
    """CSV/TSV маскируем по ячейкам: замена в одной ячейке не ломает кавычки и разделители соседних."""
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    buf = io.StringIO()
    csv.writer(buf, delimiter=delimiter, lineterminator="\n").writerows([[fn(c) for c in r] for r in rows])
    return buf.getvalue()


def _map_json(value: Any, fn: Callable[[str], str]) -> Any:
    if isinstance(value, str):
        return fn(value)
    if isinstance(value, dict):
        return {fn(k): _map_json(v, fn) for k, v in value.items()}
    if isinstance(value, list):
        return [_map_json(v, fn) for v in value]
    return value


def transform(filename: str, data: bytes, fn: Callable[[str], str]) -> bytes:
    ext = Path(filename).suffix.lower()
    if ext == ".xlsx":
        return rebuild_xlsx(data, fn)
    if ext not in TEXT_EXT:
        raise DomainError("Этот тип файла нельзя передать в песочницу.")
    text = decode_text(data)
    if ext in TABLE_EXT:
        return _mask_table(text, TABLE_EXT[ext], fn).encode("utf-8")
    if ext == ".json":
        try:
            return json.dumps(_map_json(json.loads(text), fn), ensure_ascii=False, indent=1).encode("utf-8")
        except ValueError:
            pass  # не JSON — маскируем как текст
    return fn(text).encode("utf-8")


def sanitize_upload(filename: str, data: bytes, pii: Pseudonymizer) -> bytes:
    """Загруженный пользователем файл → копия для песочницы без ПДн (UTF-8 / новая книга XLSX)."""
    return transform(filename, data, pii.mask)


def container_name(file_id: int, filename: str) -> str:
    """Имя файла в контейнере: только латиница и номер. Ни ФИО, ни кириллица туда не попадают."""
    return f"upload_{file_id}{Path(filename).suffix.lower()}"


def restore_generated(filename: str, data: bytes, pii: Pseudonymizer) -> bytes:
    """Файл, созданный кодом модели: в текстовых форматах и XLSX возвращаем настоящие имена."""
    if Path(filename).suffix.lower() not in TEXT_EXT | {".xlsx"}:
        return data  # картинки и прочее — как есть
    try:
        return transform(filename, data, pii.unmask)
    except Exception:  # noqa: BLE001 — повреждённый или неизвестный файл отдаём как есть
        log.warning("cannot restore names in generated %s file", Path(filename).suffix)
        return data
