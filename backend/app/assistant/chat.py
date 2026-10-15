"""Оркестратор одного хода диалога.

Путь запроса: вопрос → маскирование ПДн → LLM (основная / запасная) ⇄ инструменты сервисного слоя
(и Code Interpreter) → обратная замена псевдонимов → стрим пользователю.
Чтение данных выполняется сразу, изменение — только после «Подтвердить» (см. confirm_action).
"""

import asyncio
import io
import json
import logging
import mimetypes
from collections.abc import Awaitable, Callable
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from pathlib import Path
from string import Template
from typing import Any

from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.assistant import memory, tools
from app.assistant.llm.base import (
    CodeRun,
    Completed,
    LLMError,
    LLMRequest,
    LLMStreamBroken,
    TextDelta,
    ToolCall,
)
from app.assistant.llm.router import FallbackLLM, RouterStats
from app.assistant.pii import Pseudonymizer, StreamUnmasker
from app.assistant.sandbox import Sandbox, SandboxError
from app.config import get_settings
from app.errors import DomainError, Forbidden, NotFound
from app.models import ROLE_TITLES, ChatFile, Conversation, Message, PendingAction, User
from app.services import orders
from app.services.common import Actor, fmt_num
from app.services.permissions import PERM_TITLES, Perm, has_perm
from app.services.reports import week_bounds

log = logging.getLogger(__name__)

PROMPTS_DIR = Path(__file__).resolve().parent.parent / "prompts"
HISTORY_LIMIT = 20
WEEKDAYS = ["понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье"]

Emit = Callable[[dict[str, Any]], Awaitable[None]]


def load_prompt(name: str) -> str:
    return (PROMPTS_DIR / name).read_text(encoding="utf-8")


def build_instructions(
    actor: Actor, user_alias: str, facts: list[str], pii: Pseudonymizer, code_on: bool
) -> str:
    today = date.today()
    ws, we = week_bounds(today)
    perms = [PERM_TITLES[p] for p in Perm if has_perm(actor, p)]
    facts_text = "\n".join(f"- {pii.mask(f)}" for f in facts) if facts else "- (пока ничего)"
    return Template(load_prompt("system.md")).safe_substitute(
        today=today.isoformat(),
        weekday=WEEKDAYS[today.weekday()],
        week_start=ws.isoformat(),
        week_end=we.isoformat(),
        user_alias=user_alias,
        role_title=ROLE_TITLES.get(actor.role, actor.role),  # type: ignore[call-overload]
        permissions=", ".join(perms),
        facts=facts_text,
        code_interpreter="доступен" if code_on else "сейчас недоступен — графики и расчёты в Python не делай",
    )


@dataclass
class TurnState:
    text: str = ""
    model: str | None = None
    tool_calls: list[dict[str, Any]] = field(default_factory=list)
    code: list[dict[str, Any]] = field(default_factory=list)
    files: list[dict[str, Any]] = field(default_factory=list)
    actions: list[dict[str, Any]] = field(default_factory=list)
    llm_attempts: list[dict[str, Any]] = field(default_factory=list)
    status: str = "complete"
    error: str | None = None


class ChatService:
    def __init__(
        self,
        llm: FallbackLLM,
        session_factory: async_sessionmaker[AsyncSession],
        sandbox: Sandbox | None = None,
    ):
        self.llm = llm
        self.session_factory = session_factory
        self.sandbox = sandbox
        self.settings = get_settings()

    # ---------- публичный вход ----------

    async def run_turn(
        self,
        *,
        user: User,
        conversation_id: int,
        text: str,
        file_ids: list[int],
        emit: Emit,
        stop: asyncio.Event,
    ) -> int:
        """Выполняет ход и возвращает id сообщения ассистента. Никогда не бросает наружу ошибки LLM."""
        actor = Actor(user_id=user.id, full_name=user.full_name, role=user.role, via_assistant=True)
        async with self.session_factory() as session:
            conv = await self._own_conversation(session, user.id, conversation_id)
            user_msg = Message(
                conversation_id=conv.id, role="user", content=text, meta={"file_ids": file_ids}
            )
            assistant_msg = Message(conversation_id=conv.id, role="assistant", content="", status="streaming")
            session.add_all([user_msg, assistant_msg])
            if conv.title == "Новый диалог":
                conv.title = (text[:60] + "…") if len(text) > 60 else text
            conv.updated_at = datetime.now(UTC)
            await session.commit()
            if file_ids:
                await session.execute(
                    update(ChatFile)
                    .where(
                        ChatFile.id.in_(file_ids),
                        ChatFile.conversation_id == conv.id,
                        ChatFile.user_id == user.id,
                    )
                    .values(message_id=user_msg.id)
                )
                await session.commit()
            await emit(
                {"type": "start", "user_message_id": user_msg.id, "assistant_message_id": assistant_msg.id}
            )

        state = TurnState()
        work = asyncio.create_task(self._work(actor, conversation_id, assistant_msg.id, emit, state))
        stopper = asyncio.create_task(stop.wait())
        try:
            done, _ = await asyncio.wait({work, stopper}, return_when=asyncio.FIRST_COMPLETED)
            if stopper in done:
                work.cancel()
                try:
                    await work
                except (asyncio.CancelledError, Exception):  # noqa: BLE001
                    pass
                state.status = "interrupted"
                state.error = "Ответ остановлен пользователем."
            else:
                stopper.cancel()
                work.result()
        except LLMStreamBroken as exc:
            state.status = "interrupted"
            state.error = exc.user_message
        except LLMError as exc:
            log.warning("LLM failed: %s", exc)
            state.status = "error"
            state.error = exc.user_message
        except Exception:
            log.exception("chat turn failed")
            state.status = "error"
            state.error = "Что-то пошло не так при подготовке ответа. Попробуйте ещё раз."
        finally:
            await self._save(assistant_msg.id, state)
        await emit({"type": "done", "message": await self.message_dict(assistant_msg.id)})
        return assistant_msg.id

    # ---------- основной цикл ----------

    async def _work(
        self, actor: Actor, conversation_id: int, message_id: int, emit: Emit, state: TurnState
    ) -> None:
        async with self.session_factory() as session:
            conv = await session.get(Conversation, conversation_id)
            assert conv is not None
            pii = await Pseudonymizer.load(session)
            container_id = await self._prepare_sandbox(session, conv, pii)
            facts = [f.text for f in await memory.list_facts(session, actor.user_id)]
            user_alias = pii.alias_for("person", actor.user_id) or "пользователь"
            instructions = build_instructions(actor, user_alias, facts, pii, container_id is not None)
            conv_id = conv.id
            history = await self._history(session, conv_id, pii, exclude_id=message_id)

            ctx = tools.ToolContext(
                session=session,
                actor=actor,
                conversation_id=conv.id,
                pii=pii,
                assistant_message_id=message_id,
                sandbox=self.sandbox,
                container_id=container_id,
            )
            stats = RouterStats()
            state.llm_attempts = stats.attempts  # тот же список: попытки видны даже при обрыве
            items = list(history)
            unmasker = StreamUnmasker(pii)
            for _step in range(self.settings.llm_max_tool_steps):
                calls: list[ToolCall] = []
                request = LLMRequest(
                    instructions=instructions,
                    input=list(items),  # копия: следующий шаг дописывает в items
                    tools=tools.specs(),
                    code_container=container_id,
                    timeout=self.settings.llm_timeout_s,
                    temperature=self.settings.llm_temperature,
                )
                async for event in self.llm.stream(request, stats):
                    if isinstance(event, TextDelta):
                        chunk = unmasker.push(event.text)
                        if chunk:
                            state.text += chunk
                            await emit({"type": "delta", "text": chunk})
                    elif isinstance(event, ToolCall):
                        calls.append(event)
                    elif isinstance(event, CodeRun):
                        entry = {
                            "code": pii.unmask(event.code),
                            "logs": pii.unmask(event.logs)[:4000],
                            "error": event.error,
                        }
                        state.code.append(entry)
                        await emit({"type": "code", **entry})
                    elif isinstance(event, Completed):
                        if state.model != event.model:
                            state.model = event.model
                            await emit({"type": "model", "model": event.model})
                        items.extend(event.output_items)
                tail = unmasker.flush()
                if tail:
                    state.text += tail
                    await emit({"type": "delta", "text": tail})
                if not calls:
                    break
                if state.text and not state.text.endswith("\n"):
                    state.text += "\n\n"
                    await emit({"type": "delta", "text": "\n\n"})
                for call in calls:
                    await emit({"type": "tool", "name": call.name, "status": "start"})
                    result = await tools.execute(ctx, call.name, call.arguments)
                    # В логах — только имя инструмента и маскированные аргументы.
                    log.info("tool %s args=%s ok=%s", call.name, pii.mask(call.arguments)[:300], result.ok)
                    state.tool_calls.append(
                        {
                            "name": call.name,
                            "arguments": _safe_json(pii.unmask(call.arguments)),
                            "ok": result.ok,
                            "summary": result.summary,
                        }
                    )
                    await emit(
                        {
                            "type": "tool",
                            "name": call.name,
                            "status": "ok" if result.ok else "error",
                            "summary": result.summary,
                        }
                    )
                    items.append(
                        {
                            "type": "function_call_output",
                            "call_id": call.call_id,
                            "output": pii.mask_json(result.data),
                        }
                    )
                for action in ctx.actions:
                    if action not in state.actions:
                        state.actions.append(action)
                        await emit({"type": "action", "action": pii.unmask_obj(action)})
            else:
                state.text += "\n\n_Слишком много шагов, ответ остановлен. Уточните вопрос._"

            if container_id and state.code:
                await self._collect_generated(
                    session, conv_id, actor.user_id, message_id, container_id, pii, state, emit
                )

    # ---------- песочница ----------

    async def _prepare_sandbox(
        self, session: AsyncSession, conv: Conversation, pii: Pseudonymizer
    ) -> str | None:
        if self.sandbox is None:
            return None
        try:
            container_id, recreated = await self.sandbox.ensure(conv.container_id, f"candles-conv-{conv.id}")
        except SandboxError as exc:
            log.warning("sandbox unavailable: %s", exc)
            return None
        if container_id != conv.container_id:
            conv.container_id = container_id
        uploads = list(
            await session.scalars(
                select(ChatFile).where(ChatFile.conversation_id == conv.id, ChatFile.direction == "upload")
            )
        )
        for f in uploads:
            if f.container_path and not recreated:
                continue
            try:
                f.container_path = await self.sandbox.upload(
                    container_id, pii.mask(f.filename), sanitize_upload(f, pii)
                )
            except SandboxError as exc:
                log.warning("upload to sandbox failed: %s", exc)
        await session.commit()
        return container_id

    async def _collect_generated(
        self,
        session: AsyncSession,
        conv_id: int,
        user_id: int,
        message_id: int,
        container_id: str,
        pii: Pseudonymizer,
        state: TurnState,
        emit: Emit,
    ) -> None:
        assert self.sandbox is not None
        try:
            files = await self.sandbox.list_files(container_id)
        except SandboxError:
            return
        known = set(
            await session.scalars(
                select(ChatFile.container_path).where(
                    ChatFile.conversation_id == conv_id, ChatFile.container_path.is_not(None)
                )
            )
        )
        new_files = [
            f
            for f in files
            if f.source != "user" and f.path not in known and (f.size or 0) <= 20 * 1024 * 1024
        ]
        # Code Interpreter перехватывает незакрытую фигуру как cfile_*.png. Если модель сохранила
        # картинку и сама, перехваченная копия — дубль.
        if any(_is_image(f.path) and not _captured(f.path) for f in new_files):
            new_files = [f for f in new_files if not (_is_image(f.path) and _captured(f.path))]
        for f in new_files:
            try:
                data = await self.sandbox.download(container_id, f.id)
            except SandboxError:
                continue
            filename = f.path.rsplit("/", 1)[-1]
            if is_chart_spec(filename, data):
                continue  # служебная заготовка интерактивного графика Code Interpreter, пользователю не нужна
            mime = mimetypes.guess_type(filename)[0] or "application/octet-stream"
            if filename.startswith("cfile_") and mime.startswith("image/"):
                # Картинка, перехваченная из plt.show(): даём человеческое имя вместо технического.
                n = sum(1 for x in state.files if x["mime"].startswith("image/")) + 1
                filename = f"график_{n}{Path(filename).suffix}"
            data = restore_generated(filename, data, pii)
            row = ChatFile(
                conversation_id=conv_id,
                message_id=message_id,
                user_id=user_id,
                direction="generated",
                filename=filename,
                mime=mime,
                data=data,
                container_path=f.path,
            )
            session.add(row)
            await session.commit()
            info = {"id": row.id, "filename": filename, "mime": mime}
            state.files.append(info)
            await emit({"type": "file", "file": info})

    # ---------- история, сохранение ----------

    async def _history(
        self, session: AsyncSession, conv_id: int, pii: Pseudonymizer, exclude_id: int
    ) -> list[dict[str, Any]]:
        rows = list(
            await session.scalars(
                select(Message)
                .where(Message.conversation_id == conv_id, Message.id != exclude_id)
                .order_by(Message.id.desc())
                .limit(HISTORY_LIMIT)
            )
        )[::-1]
        uploaded = {
            f.id: f
            for f in await session.scalars(
                select(ChatFile).where(ChatFile.conversation_id == conv_id, ChatFile.direction == "upload")
            )
        }
        items: list[dict[str, Any]] = []
        for m in rows:
            content = m.content
            if m.role == "user":
                paths = [uploaded[i].container_path for i in m.meta.get("file_ids", []) if i in uploaded]
                if paths:
                    content += "\n\n[Пользователь загрузил файлы: " + ", ".join(p for p in paths if p) + "]"
            elif m.status == "interrupted":
                content += "\n[ответ был прерван]"
            if not content.strip():
                continue
            items.append({"role": m.role, "content": pii.mask(content)})
        return items

    async def _save(self, message_id: int, state: TurnState) -> None:
        async with self.session_factory() as session:
            msg = await session.get(Message, message_id)
            if msg is None:
                return
            msg.content = state.text
            msg.status = state.status
            msg.model = state.model
            msg.meta = {
                "tool_calls": state.tool_calls,
                "code": state.code,
                "files": state.files,
                "action_ids": [a["id"] for a in state.actions],
                "error": state.error,
                "llm_attempts": state.llm_attempts,
            }
            await session.commit()

    async def message_dict(self, message_id: int) -> dict[str, Any]:
        async with self.session_factory() as session:
            msg = await session.get(Message, message_id)
            assert msg is not None
            return await message_to_dict(session, msg)

    async def _own_conversation(
        self, session: AsyncSession, user_id: int, conversation_id: int
    ) -> Conversation:
        conv = await session.get(Conversation, conversation_id)
        if conv is None or conv.user_id != user_id:
            raise NotFound("Диалог не найден.")
        return conv


# ---------- подтверждение действий ----------


async def confirm_action(session: AsyncSession, user: User, action_id: int) -> PendingAction:
    """Пользователь нажал «Подтвердить». Вызывается тот же сервис, что и у кнопки «Создать задание».

    Строка действия заблокирована (FOR UPDATE) до конца транзакции, а задание создаётся и действие
    помечается подтверждённым одним коммитом. Поэтому двойной клик не создаст два задания:
    второй запрос дождётся блокировки и увидит статус «confirmed».
    """
    action = await session.scalar(
        select(PendingAction).where(PendingAction.id == action_id).with_for_update()
    )
    if action is None or action.user_id != user.id:
        raise NotFound("Действие не найдено.")
    if action.status != "pending":
        raise DomainError("Это действие уже обработано.")
    actor = Actor(user_id=user.id, full_name=user.full_name, role=user.role, via_assistant=True)
    p = action.payload
    try:
        if action.kind != "create_and_plan_order":
            raise DomainError("Неизвестное действие.")
        result = await orders.create_and_plan(
            session,
            actor,
            recipe_id=p["recipe_id"],
            qty=p["qty"],
            planned_date=date.fromisoformat(p["planned_date"]),
            workshop=p["workshop"],
            comment=p.get("comment"),
            commit=False,
        )
    except (DomainError, Forbidden) as exc:
        await session.rollback()
        action = await session.get(PendingAction, action_id)
        assert action is not None
        action.status = "failed"
        action.result = {"error": exc.message}
        action.decided_at = datetime.now(UTC)
        session.add(
            Message(
                conversation_id=action.conversation_id,
                role="assistant",
                content=f"Не удалось выполнить действие: {exc.message}",
                meta={"system_note": True},
            )
        )
        await session.commit()
        return action

    order = result.order
    shortage_text = "; ".join(f"{fmt_num(s.missing)} {s.unit} «{s.material_name}»" for s in result.shortages)
    action.status = "confirmed"
    action.decided_at = datetime.now(UTC)
    action.result = {"order_id": order.id, "order_number": order.number, "shortage": shortage_text or None}
    note = f"Пользователь подтвердил действие. Создано и запланировано задание {order.number}."
    note += f" Дефицит: {shortage_text}." if shortage_text else " Всё зарезервировано."
    session.add(
        Message(
            conversation_id=action.conversation_id, role="assistant", content=note, meta={"system_note": True}
        )
    )
    await session.commit()
    return action


async def reject_action(session: AsyncSession, user: User, action_id: int) -> PendingAction:
    action = await session.scalar(
        select(PendingAction).where(PendingAction.id == action_id).with_for_update()
    )
    if action is None or action.user_id != user.id:
        raise NotFound("Действие не найдено.")
    if action.status != "pending":
        raise DomainError("Это действие уже обработано.")
    action.status = "rejected"
    action.decided_at = datetime.now(UTC)
    session.add(
        Message(
            conversation_id=action.conversation_id,
            role="assistant",
            content="Пользователь отклонил предложенное действие.",
            meta={"system_note": True},
        )
    )
    await session.commit()
    return action


# ---------- сериализация ----------


async def message_to_dict(session: AsyncSession, msg: Message) -> dict[str, Any]:
    actions = []
    ids = msg.meta.get("action_ids") or []
    if ids:
        for a in await session.scalars(
            select(PendingAction).where(PendingAction.id.in_(ids)).order_by(PendingAction.id)
        ):
            actions.append(
                {
                    "id": a.id,
                    "kind": a.kind,
                    "status": a.status,
                    "preview": a.preview,
                    "result": a.result,
                    "decided_at": a.decided_at.isoformat() if a.decided_at else None,
                }
            )
    files = []
    if msg.role == "user" and msg.meta.get("file_ids"):
        for f in await session.scalars(select(ChatFile).where(ChatFile.id.in_(msg.meta["file_ids"]))):
            files.append({"id": f.id, "filename": f.filename, "mime": f.mime})
    else:
        files = msg.meta.get("files") or []
    return {
        "id": msg.id,
        "role": msg.role,
        "content": msg.content,
        "status": msg.status,
        "model": msg.model,
        "tool_calls": msg.meta.get("tool_calls") or [],
        "code": msg.meta.get("code") or [],
        "files": files,
        "actions": actions,
        "error": msg.meta.get("error"),
        "fallback_used": any(not a.get("ok") for a in msg.meta.get("llm_attempts") or []),
        "llm_attempts": msg.meta.get("llm_attempts") or [],
        "system_note": bool(msg.meta.get("system_note")),
        "created_at": msg.created_at.isoformat() if msg.created_at else None,
    }


def _safe_json(raw: str) -> Any:
    try:
        return json.loads(raw)
    except ValueError:
        return raw[:500]


# ---------- файлы ----------

TEXT_EXT = {".csv", ".txt", ".tsv", ".json", ".md"}


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _map_xlsx(data: bytes, fn: Callable[[str], str]) -> bytes:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = fn(cell.value)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def sanitize_upload(f: ChatFile, pii: Pseudonymizer) -> bytes:
    """Загруженный пользователем файл перед отправкой в песочницу: ПДн → псевдонимы."""
    ext = Path(f.filename).suffix.lower()
    if ext in TEXT_EXT:
        return pii.mask(_decode(f.data)).encode("utf-8")
    if ext == ".xlsx":
        return _map_xlsx(f.data, pii.mask)
    return f.data


def _is_image(path: str) -> bool:
    return (mimetypes.guess_type(path)[0] or "").startswith("image/")


def _captured(path: str) -> bool:
    return path.rsplit("/", 1)[-1].startswith("cfile_")


def is_chart_spec(filename: str, data: bytes) -> bool:
    """JSON-описание графика, которое Code Interpreter создаёт рядом с картинкой ({type, labels, datasets})."""
    if not filename.endswith(".json"):
        return False
    try:
        spec = json.loads(data)
    except ValueError:
        return False
    return isinstance(spec, dict) and {"labels", "datasets"} <= spec.keys()


def restore_generated(filename: str, data: bytes, pii: Pseudonymizer) -> bytes:
    """Файл, созданный кодом модели: в текстовых форматах возвращаем настоящие имена."""
    ext = Path(filename).suffix.lower()
    try:
        if ext in TEXT_EXT:
            return pii.unmask(_decode(data)).encode("utf-8")
        if ext == ".xlsx":
            return _map_xlsx(data, pii.unmask)
    except Exception:  # noqa: BLE001 — повреждённый файл отдаём как есть
        log.warning("cannot restore names in %s", filename)
    return data
